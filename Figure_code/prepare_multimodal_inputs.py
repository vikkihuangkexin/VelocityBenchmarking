#!/usr/bin/env python3
"""Import ATAC/Spatial matrices and extract standalone resource/usability inputs.

Normal analysis uses the distributed CSV files and does not need the archives or
workbook. This optional preparation script documents the extraction exactly.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import math
from pathlib import Path
import re
import zipfile

import pandas as pd


ATAC_METHODS = ['MoFlow', 'ArchVelo', 'DynaVelo', 'CRAK-Velo', 'MultiVeloVAE',
                'GraphVelo', 'InterVelo', 'MultiVelo', 'scKINETICS']
SPATIAL_METHODS = ['TopoVelo', 'STEER', 'spaVelo', 'STT', 'SDEvelo', 'cell2fate', 'veloAgent']
SPATIAL_CONTROLS = ['InterVelo', 'scVelo-dynamical', 'scVelo-stochastic']
METRICS = {'AngleConsistency': 'Angle consistency', 'CBDir': 'CBDir',
           'TransitionScore': 'Transition V1', 'GroundTruthCorrelation': 'Ground-truth correlation',
           'ConsistencyScore': 'Consistency score', 'ICCoh': 'ICCoh', 'STC': 'STC', 'SVC': 'SVC'}


def key(value):
    return re.sub('[^a-z0-9]', '', str(value).lower())


def sha(data):
    return hashlib.sha256(data).hexdigest()


def english(text):
    if any('\u3400' <= c <= '\u9fff' for c in str(text)):
        raise ValueError('CJK text is not allowed in distributed filenames or text data')
    return text


def save(frame, path):
    path.parent.mkdir(parents=True, exist_ok=True)
    data = frame.to_csv(index=False, float_format='%.17g')
    english(data)
    path.write_text(data, encoding='utf-8')


def import_archives(atac, spatial, root):
    provenance = []
    for profile, archive in [('ATAC', atac), ('Spatial', spatial)]:
        with zipfile.ZipFile(archive) as z:
            for member in z.namelist():
                name = Path(member).name
                if not name.endswith('.csv'):
                    continue
                if profile == 'ATAC':
                    match = re.fullmatch(r'(\w+)_ATAC_(real|sim)_02_summary.csv', name)
                    if not match:
                        raise ValueError(f'Unrecognized ATAC summary: {name}')
                    prefix, typ = match.groups()
                    typ = 'simulation' if typ == 'sim' else 'real'
                    folder = 'summary'
                elif name.startswith('spatial_time_consistency_'):
                    prefix = 'STC'
                    typ = 'real' if name.endswith('_real.csv') else 'downsampling'
                    folder = 'summary' if typ == 'real' else 'downsampling'
                elif name.startswith('spatial_velocity_consistency_'):
                    prefix = 'SVC'
                    typ = 'real' if name.endswith('_real.csv') else 'downsampling'
                    folder = 'summary' if typ == 'real' else 'downsampling'
                else:
                    match = re.fullmatch(r'(\w+)_(?:spatail|spatial)_(real|simulation)_02_summary.csv', name)
                    if not match:
                        raise ValueError(f'Unrecognized Spatial summary: {name}')
                    prefix, typ = match.groups()
                    folder = 'summary' if typ == 'real' else 'downsampling'
                    if typ == 'simulation':
                        typ = 'downsampling'
                if profile == 'Spatial' and prefix not in {'AngleConsistency', 'CBDir', 'TransitionScore', 'STC', 'SVC'}:
                    continue
                if prefix not in METRICS:
                    raise ValueError(f'Unrecognized metric: {prefix}')
                data = z.read(member)
                english(data.decode('utf-8-sig'))
                # Preserve supplied matrix bytes. Only its path is standardized.
                path = root / f'measurements_raw/{profile}/measurement/{folder}/{prefix}_{profile}_{typ}_02_summary.csv'
                path.parent.mkdir(parents=True, exist_ok=True)
                path.write_bytes(data)
                provenance.append({'profile': profile, 'source_kind': 'uploaded_archive',
                    'source_name': english(archive.name), 'source_sha256': sha(archive.read_bytes()),
                    'source_member': english(name), 'source_member_sha256': sha(data),
                    'output_file': path.relative_to(root).as_posix(),
                    'note': 'Matrix bytes preserved; Spatial perturbations routed to downsampling only'})
    for profile, methods in [('ATAC', ATAC_METHODS), ('Spatial', SPATIAL_METHODS)]:
        save(pd.DataFrame({'method_order': range(1, len(methods) + 1), 'method': methods}),
             root / f'measurements_raw/{profile}/methods.csv')
    return provenance


def cost(value):
    token = '' if value is None else str(value)
    if token.strip().lower() in {'', 'na', 'nan', 'n/a', 'error', 'none'}:
        return token, math.nan, 'missing_marker'
    try:
        number = float(value)
    except (ValueError, TypeError):
        english(token)
        return token, math.nan, 'recorded_run_failure'
    if not math.isfinite(number) or number < 0:
        raise ValueError('Invalid resource measurement')
    return token, number, 'valid_zero' if number == 0 else 'valid_positive'


def extract_workbook(workbook, root):
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
    wb = load_workbook(workbook, read_only=True, data_only=True)
    source_hash = sha(workbook.read_bytes())
    provenance = []
    dictionary = []
    for row in wb['Table S3'].iter_rows(max_col=3, values_only=True):
        if row[0] is not None and str(row[0]).strip().isdigit() and row[2]:
            dictionary.append({'dataset_number': str(row[0]), 'dataset_id': str(row[2])})
    for profile, bounds in [('ATAC', (73, 91)), ('Spatial', (55, 72))]:
        frame = pd.DataFrame([r for r in dictionary if bounds[0] <= int(r['dataset_number']) <= bounds[1]])
        save(frame, root / f'measurements_raw/{profile}/real_dataset_ids.csv')
    # Resource blocks are identified from their headers, not fixed row numbers.
    for sheet, measure, unit, suffix in [('Table S4', 'runtime_seconds', 'seconds', 'runtime'),
                                         ('Table S5', 'ram_gb', 'GB', 'cpu_ram'),
                                         ('Table S6', 'vram_gb', 'GB', 'gpu_vram')]:
        rows = list(wb[sheet].iter_rows(max_col=75, values_only=True))
        for profile, methods in [('ATAC', ATAC_METHODS), ('Spatial', SPATIAL_METHODS)]:
            names = {key(m): m for m in methods}
            records = []
            for hi, header in enumerate(rows):
                if key(header[0]) != 'method':
                    continue
                cols = [(ci, str(v).strip()) for ci, v in enumerate(header) if ci and v is not None]
                if profile == 'ATAC':
                    cols = [(ci, d) for ci, d in cols if re.fullmatch(r'\d+_.+_c\d+_g\d+', d)]
                else:
                    cols = [(ci, d) for ci, d in cols if re.fullmatch(r'[\dKk]+\s*[Xx]\s*[\dKk]+', d)]
                if not cols:
                    continue
                end = next((j for j in range(hi + 1, len(rows)) if key(rows[j][0]) == 'method'), len(rows))
                seen = {}
                for ri in range(hi + 1, end):
                    method = names.get(key(rows[ri][0]))
                    if method:
                        if method in seen:
                            raise ValueError(f'Duplicate resource method: {sheet}/{profile}/{method}')
                        seen[method] = ri
                for ci, dataset in cols:
                    if profile == 'ATAC':
                        match = re.fullmatch(r'(.+)_c(\d+)_g(\d+)', dataset)
                        ds = f'{match[1]}_cell{match[2]}_gene{match[3]}'
                        cells, genes = int(match[2]), int(match[3])
                    else:
                        def size(t):
                            return int(float(t[:-1]) * 1000) if t.upper().endswith('K') else int(t)
                        a, b = re.split(r'\s*[Xx]\s*', dataset)
                        cells, genes = size(a), size(b)
                        ds = f'cell_scaling_cell{cells}_gene{genes}' if genes == 9836 else f'gene_scaling_cell{cells}_gene{genes}'
                    for method in methods:
                        ri = seen.get(method)
                        raw = rows[ri][ci] if ri is not None else None
                        token, value, status = cost(raw)
                        if ri is None:
                            status = 'method_not_provided'
                        records.append({'method': method, 'protocol': 'top_' + profile,
                            'measure': measure, 'dataset_id': ds, 'unit': unit,
                            'clean_value': value, 'raw_token': token, 'value_status': status,
                            'cells': cells, 'genes': genes, 'data_type': 'simulation' if profile == 'ATAC' else 'real',
                            'source_locator': f'{sheet}!{get_column_letter(ci + 1)}{ri + 1}' if ri is not None else '',
                            'source_workbook': workbook.name})
            path = root / f'scalability/{profile}_{suffix}.csv'
            save(pd.DataFrame(records), path)
            provenance.append({'profile': profile, 'source_kind': 'supplementary_workbook',
                'source_name': workbook.name, 'source_sha256': source_hash, 'source_member': sheet,
                'source_member_sha256': '', 'output_file': path.relative_to(root).as_posix(),
                'note': '21 ATAC workloads or 11 Spatial workloads; source cells retained'})
    for profile, methods in [('ATAC', ATAC_METHODS), ('Spatial', SPATIAL_METHODS)]:
        names = {key(m): m for m in methods}
        items = []
        for ri, row in enumerate(wb['Table S8'].iter_rows(max_col=9, values_only=True), 1):
            method = names.get(key(row[0]))
            if method:
                category = {'Install-friendly': 'Installability', 'Paper': 'Publication quality'}.get(row[1], row[1])
                items.append({'method': method, 'category': category, 'category_weight': row[2],
                    'reported_category_score': row[3], 'subcategory': row[4], 'subscore': row[5],
                    'subcategory_weight': row[6], 'source_locator': f'Table S8!A{ri}:I{ri}',
                    'evidence_note': row[7] or '', 'source_url': row[8] or ''})
        if len(items) != len(methods) * 24:
            raise ValueError(f'Incomplete usability items: {profile}')
        path = root / f'usability/{profile}_usability_items.csv'
        save(pd.DataFrame(items), path)
        provenance.append({'profile': profile, 'source_kind': 'supplementary_workbook',
            'source_name': workbook.name, 'source_sha256': source_hash, 'source_member': 'Table S8',
            'source_member_sha256': '', 'output_file': path.relative_to(root).as_posix(),
            'note': '24 item scores per method; category totals recomputed by analysis'})
    wb.close()
    return provenance


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--atac-zip', type=Path, required=True)
    parser.add_argument('--spatial-zip', type=Path, required=True)
    parser.add_argument('--workbook', type=Path, required=True)
    parser.add_argument('--source-root', type=Path, default=Path(__file__).resolve().parents[1] / 'Source_data')
    args = parser.parse_args()
    rows = import_archives(args.atac_zip, args.spatial_zip, args.source_root)
    rows += extract_workbook(args.workbook, args.source_root)
    save(pd.DataFrame(rows), args.source_root / 'multimodal_sources.csv')
    print(f'Prepared {len(rows)} multimodal input files')


if __name__ == '__main__':
    main()
